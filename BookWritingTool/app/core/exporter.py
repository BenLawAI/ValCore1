"""
Exporter - Export manuscripts to Word, PDF, and plain text formats
"""
from pathlib import Path
from typing import List, Optional
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.units import inch
from loguru import logger
import markdown


class Exporter:
    """Export manuscripts to various formats"""

    def __init__(self):
        """Initialize exporter"""
        pass

    def export_to_word(
        self,
        chapters: List[Path],
        output_file: Path,
        title: str = "Untitled",
        author: str = "Author",
        include_headers: bool = True,
    ) -> Path:
        """Export manuscript to Microsoft Word (.docx)

        Args:
            chapters: List of chapter markdown files
            output_file: Output .docx file path
            title: Book title
            author: Author name
            include_headers: Include chapter headers

        Returns:
            Path to created Word document
        """
        logger.info(f"Exporting to Word: {output_file}")

        doc = Document()

        # Set default font
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)

        # Title page
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run(title)
        title_run.font.size = Pt(24)
        title_run.font.bold = True

        doc.add_paragraph()

        author_para = doc.add_paragraph()
        author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_run = author_para.add_run(f"by {author}")
        author_run.font.size = Pt(14)

        doc.add_page_break()

        # Chapters
        for idx, chapter_file in enumerate(chapters, start=1):
            if not chapter_file.exists():
                logger.warning(f"Chapter not found: {chapter_file}")
                continue

            # Chapter header
            if include_headers:
                chapter_title = chapter_file.stem.replace("_", " ").title()
                heading = doc.add_heading(f"Chapter {idx}: {chapter_title}", level=1)
                doc.add_paragraph()

            # Chapter content
            content = chapter_file.read_text(encoding="utf-8")

            # Convert markdown to paragraphs (simple conversion)
            paragraphs = content.split("\n\n")

            for para_text in paragraphs:
                para_text = para_text.strip()
                if para_text:
                    # Handle markdown headings
                    if para_text.startswith("# "):
                        doc.add_heading(para_text[2:], level=2)
                    elif para_text.startswith("## "):
                        doc.add_heading(para_text[3:], level=3)
                    else:
                        # Regular paragraph
                        para = doc.add_paragraph(para_text)
                        para.paragraph_format.first_line_indent = Inches(0.5)
                        para.paragraph_format.line_spacing = 2.0  # Double-spaced

            # Page break between chapters
            if idx < len(chapters):
                doc.add_page_break()

        # Save document
        output_file.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_file))

        logger.info(f"Word export complete: {output_file}")
        return output_file

    def export_to_pdf(
        self,
        chapters: List[Path],
        output_file: Path,
        title: str = "Untitled",
        author: str = "Author",
        include_headers: bool = True,
    ) -> Path:
        """Export manuscript to PDF

        Args:
            chapters: List of chapter markdown files
            output_file: Output .pdf file path
            title: Book title
            author: Author name
            include_headers: Include chapter headers

        Returns:
            Path to created PDF
        """
        logger.info(f"Exporting to PDF: {output_file}")

        output_file.parent.mkdir(parents=True, exist_ok=True)

        # Create PDF
        doc = SimpleDocTemplate(
            str(output_file),
            pagesize=letter,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=72,
        )

        # Styles
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Heading1"],
            fontSize=24,
            textColor="black",
            spaceAfter=30,
            alignment=1,  # Center
        )

        author_style = ParagraphStyle("Author", parent=styles["Normal"], fontSize=14, alignment=1, spaceAfter=50)

        chapter_heading_style = ParagraphStyle(
            "ChapterHeading", parent=styles["Heading1"], fontSize=18, spaceAfter=12
        )

        body_style = ParagraphStyle(
            "Body", parent=styles["Normal"], fontSize=12, leading=24, firstLineIndent=0.5 * inch
        )

        # Build story
        story = []

        # Title page
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 0.2 * inch))
        story.append(Paragraph(f"by {author}", author_style))
        story.append(PageBreak())

        # Chapters
        for idx, chapter_file in enumerate(chapters, start=1):
            if not chapter_file.exists():
                logger.warning(f"Chapter not found: {chapter_file}")
                continue

            # Chapter header
            if include_headers:
                chapter_title = chapter_file.stem.replace("_", " ").title()
                story.append(Paragraph(f"Chapter {idx}: {chapter_title}", chapter_heading_style))
                story.append(Spacer(1, 0.3 * inch))

            # Chapter content
            content = chapter_file.read_text(encoding="utf-8")
            paragraphs = content.split("\n\n")

            for para_text in paragraphs:
                para_text = para_text.strip()
                if para_text and not para_text.startswith("#"):
                    # Clean markdown formatting (simple)
                    para_text = para_text.replace("**", "").replace("*", "").replace("_", "")
                    story.append(Paragraph(para_text, body_style))
                    story.append(Spacer(1, 0.15 * inch))

            # Page break between chapters
            if idx < len(chapters):
                story.append(PageBreak())

        # Build PDF
        doc.build(story)

        logger.info(f"PDF export complete: {output_file}")
        return output_file

    def export_to_text(
        self, chapters: List[Path], output_file: Path, include_headers: bool = True
    ) -> Path:
        """Export manuscript to plain text

        Args:
            chapters: List of chapter markdown files
            output_file: Output .txt file path
            include_headers: Include chapter headers

        Returns:
            Path to created text file
        """
        logger.info(f"Exporting to plain text: {output_file}")

        output_file.parent.mkdir(parents=True, exist_ok=True)

        lines = []

        for idx, chapter_file in enumerate(chapters, start=1):
            if not chapter_file.exists():
                logger.warning(f"Chapter not found: {chapter_file}")
                continue

            if include_headers:
                chapter_title = chapter_file.stem.replace("_", " ").title()
                lines.append(f"\nChapter {idx}: {chapter_title}\n")
                lines.append("=" * 60)
                lines.append("\n")

            content = chapter_file.read_text(encoding="utf-8")
            lines.append(content)
            lines.append("\n\n")

        # Write file
        final_text = "\n".join(lines)
        output_file.write_text(final_text, encoding="utf-8")

        logger.info(f"Text export complete: {output_file}")
        return output_file

    def get_word_count(self, chapters: List[Path]) -> int:
        """Get total word count from chapters

        Args:
            chapters: List of chapter files

        Returns:
            Total word count
        """
        total_words = 0

        for chapter_file in chapters:
            if chapter_file.exists():
                content = chapter_file.read_text(encoding="utf-8")
                total_words += len(content.split())

        return total_words
